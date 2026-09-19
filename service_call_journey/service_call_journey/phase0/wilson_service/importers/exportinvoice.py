"""ExportInvoice (Invoice Maintenance) importer — the full open-ticket list with every ePASS status.

xlsx, sheet 'Invoice Maintenance', header on row 3. `Invoice #` is left-padded with spaces.
This file is the only place SO2/SO2.2/SO3/SO4/SO7/SO8/SI*/WAR* show up until ePASS widens
the DispatchTrack feed, so it is what confirms most keyed sync items.

`* Sched Date` is unreliable for unscheduled tickets (ePASS puts a far-future placeholder);
it only becomes epass_route_date when it is within `import.ei_sched_horizon_days` and the
job is not currently in the DispatchTrack feed (DT is authoritative when present).
"""
from __future__ import annotations

import dataclasses
import datetime as _dt
import os
from typing import Optional

from ..db import DB, now_iso
from .. import seed, sync, stuck, kpi
from .common import (canonical_code, compact_json, ensure_zone, find_or_create_address, find_or_create_customer, parse_date, parse_money,
                     split_name, tech_id_for, upsert_unit, warranty_flags, zone_for_zip, any_work_days, is_parking_day)

HEADER_ROW = 3
REQUIRED = ("Job Status", "Invoice #", "Name", "SP")


class AlreadyImported(Exception):
    pass


@dataclasses.dataclass
class ImportResult:
    import_batch_id: int
    file_name: str
    rows: int = 0
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    attached: int = 0
    new_statuses: int = 0
    sync: dict = dataclasses.field(default_factory=dict)
    stale: int = 0
    parked: int = 0            # 9/15: sched date on a day nobody works = ePASS's parking date, not a schedule

    def summary(self) -> str:
        return (f"{self.file_name}: {self.rows} tickets -> created {self.created}, updated {self.updated}, unchanged {self.unchanged}, "
                f"attached {self.attached}, new statuses {self.new_statuses}, parked {self.parked}, sync {self.sync}, stale {self.stale}")


def read_rows(path: str) -> list[dict]:
    import openpyxl  # stdlib-free fallback is not worth it; openpyxl ships with the sales-side importer
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < HEADER_ROW:
        raise ValueError("ExportInvoice: fewer than 3 rows")
    hdr = [str(c).strip() if c is not None else "" for c in rows[HEADER_ROW - 1]]
    missing = [c for c in REQUIRED if c not in hdr]
    if missing:
        raise ValueError(f"ExportInvoice: header row {HEADER_ROW} missing columns {missing}; got {hdr}")
    out = []
    for r in rows[HEADER_ROW:]:
        d = dict(zip(hdr, r))
        inv = str(d.get("Invoice #") or "").strip()
        if not inv:
            continue
        d["Invoice #"] = inv
        out.append({k: (v.strip() if isinstance(v, str) else v) for k, v in d.items() if k})
    return out


def import_file(db: DB, path: str, *, now: Optional[_dt.datetime] = None, file_name: Optional[str] = None) -> ImportResult:
    now = now or _dt.datetime.now()
    ts = now_iso(now)
    file_name = file_name or os.path.basename(path)
    if db.fetchone("SELECT 1 FROM import_batch WHERE file_name=?", (file_name,)):
        raise AlreadyImported(file_name)
    try:
        mtime = _dt.datetime.fromtimestamp(os.path.getmtime(path)).replace(microsecond=0).isoformat(sep=" ")
    except OSError:
        mtime = None
    batch_id = db.insert("import_batch", {"source": "EI", "file_name": file_name, "file_path": path, "file_modified_at": mtime, "imported_at": ts,
                                          "row_count": 0, "sv_count": 0, "created_count": 0, "updated_count": 0, "status": "running", "message": None})
    db.commit()
    res = ImportResult(import_batch_id=batch_id, file_name=file_name)
    try:
        rows = read_rows(path)
        res.rows = len(rows)
        horizon = db.setting("import.ei_sched_horizon_days", 120)
        workdays = any_work_days(db)
        seen: dict[str, int] = {}
        for i, r in enumerate(rows, 1):
            sv = r["Invoice #"]
            db.insert("import_row_raw", {"import_batch_id": batch_id, "order_number": sv, "line_no": i, "job_status": r.get("Job Status"),
                                         "delivery_date": parse_date(r.get("* Sched Date")), "truck": r.get("SP") or r.get("Route"), "map_zone": r.get("Map Zone"),
                                         "model": r.get("Service Model"), "description": None, "quantity": None, "amount": parse_money(r.get("Total")), "row_json": compact_json(r)})
            job_id, outcome, new_status, parked = _upsert_ticket(db, r, ts, now, horizon, workdays)
            seen[sv] = job_id
            res.new_statuses += new_status
            res.parked += parked
            setattr(res, outcome, getattr(res, outcome) + 1)
        res.sync = sync.reconcile_after_import(db, batch_id, seen, source="EI", now=now)
        res.stale = stuck.mark_stale(db, now)
        db.update("import_batch", {"import_batch_id": batch_id}, {"row_count": res.rows, "sv_count": res.rows, "created_count": res.created,
                                                                  "updated_count": res.updated, "status": "ok", "message": res.summary()})
        db.commit()
        return res
    except Exception as e:
        db.rollback()
        db.update("import_batch", {"import_batch_id": batch_id}, {"status": "failed", "message": f"{type(e).__name__}: {e}"[:2000]})
        db.commit()
        raise


def _upsert_ticket(db: DB, r: dict, ts: str, now: _dt.datetime, horizon_days: int, workdays: set) -> tuple[int, str, int, int]:
    sv = r["Invoice #"]
    epass_status = (str(r.get("Job Status") or "").strip() or None)
    if epass_status:
        epass_status = epass_status.upper()   # 9/15: the real export carries both SO8 and so8
    new_status = 1 if seed.ensure_status(db, epass_status) else 0
    sp = canonical_code(db, r.get("SP") or r.get("Route"))
    tech_id = tech_id_for(db, sp)
    zone_code = ensure_zone(db, r.get("Map Zone")) or zone_for_zip(db, r.get("Zip Code"))
    flags, is_wty = warranty_flags(r.get("Priorities"))
    pay = str(r.get("Payment Type Code") or "").strip()[:4] or None
    qual = str(r.get("Qualification") or "").strip()[:6] or None
    balance, total = parse_money(r.get("Balance")), parse_money(r.get("Total"))
    finish, created = parse_date(r.get("Finish Date")), parse_date(r.get("Date Created"))
    sched = parse_date(r.get("* Sched Date"))
    parked = 0
    if sched and (_dt.date.fromisoformat(sched) - now.date()).days > horizon_days:
        sched = None  # far-future placeholder = unscheduled
    if is_parking_day(db, sched, workdays):
        sched, parked = None, 1
    inv_status = str(r.get("Status") or "").strip()[:12] or None
    brand = str(r.get("Service Brand") or "").strip()[:40] or None
    model = str(r.get("Service Model") or "").strip()[:60] or None
    serial = str(r.get("Service Serial") or "").strip()[:60] or None

    job = db.fetchone("SELECT * FROM job WHERE sv_number=?", (sv,))
    attached = False
    if job is None:
        job = sync.match_create_ticket(db, sv, phone=r.get("Bill To Customer"), last_name=r.get("Name"), zip_code=r.get("Zip Code"), now=ts)
        attached = job is not None

    mirror = {"epass_status": epass_status, "epass_tech_code": sp, "epass_invoice_status": inv_status, "epass_finish_date": finish,
              "epass_created_at": created, "balance": balance, "total": total, "warranty_flags": flags, "is_warranty": is_wty,
              "payment_type": pay, "qualification": qual, "units": int(r["Units"]) if isinstance(r.get("Units"), (int, float)) else None}

    if job is None:
        first, last, disp = split_name(_last_first_to_first_last(r.get("Name")))
        code = str(r.get("Bill To Customer") or "").strip() or None
        cust_id = find_or_create_customer(db, name=disp, phone=code if code and code.isdigit() else None, alt_phone=None,
                                          email=r.get("Bill To Email"), epass_code=code, now=ts)
        addr_id = find_or_create_address(db, cust_id, line1=r.get("Address"), line2=None, city=None, state="TX", zip_code=r.get("Zip Code"),
                                         lat=None, lng=None, directions=None, zone_code=zone_code)
        job_type = "hvac" if qual == "HVAC" or (brand or "").upper() in ("TRANE", "CARR", "CARRIER", "GOODM", "LENNOX") else "appliance"
        vals = dict(mirror)
        vals.update({
            "sv_number": sv, "customer_id": cust_id, "address_id": addr_id, "status": epass_status, "status_changed_at": ts, "flags": None,
            "job_type": job_type, "source": "import", "owner_tech_id": tech_id if epass_status and epass_status not in ("SO1", "SO1.AUTH", "REQ") else None,
            "assigned_tech_id": tech_id, "promised_window_start": None, "promised_window_end": None, "planned_slot_start": None, "planned_slot_end": None,
            "route_date": sched, "route_sequence": None, "route_locked": 0, "trip_id": None,
            "booking_mode": (db.fetchone("SELECT booking_mode FROM zone WHERE zone_code=?", (zone_code,)) or {}).get("booking_mode") if zone_code else None,
            "zone_code": zone_code, "problem_text": None, "bin_location": None,
            "epass_route_date": sched, "epass_seen_at": ts, "epass_source": "EI", "in_feed": 0,
            "stale": 0, "needs_intake_review": 1, "closed_at": None, "cancel_reason": None, "created_at": ts, "updated_at": ts,
        })
        job_id = db.insert("job", vals)
        db.insert("status_history", {"job_id": job_id, "from_status": None, "to_status": epass_status, "changed_at": ts, "actor_type": "import",
                                     "actor_id": "EI", "trigger_event": "import.create", "reason_code": None, "note": f"created from ExportInvoice; SP {sp or '-'}"})
        upsert_unit(db, job_id, {"brand": brand, "model": model, "serial": serial}, model=model, serial=serial, brand=brand)
        if parked:
            db.update("job", {"job_id": job_id}, {"flags": "parked"})
        kpi.detect_recall(db, job_id, now)
        return job_id, "created", new_status, parked

    changes = {k: v for k, v in mirror.items() if v is not None and job.get(k) != v}
    if job.get("in_feed"):
        changes.pop("epass_tech_code", None)  # DispatchTrack's Truck is the routing authority while the job is in its feed
    # DT is authoritative for route date while the job is in its feed
    if not job.get("in_feed") and sched != job.get("epass_route_date"):
        changes["epass_route_date"] = sched
    if not job.get("in_feed") and job.get("epass_source") != "EI":
        changes["epass_source"] = "EI"
    if job.get("zone_code") is None and zone_code:
        changes["zone_code"] = zone_code
    substantive = set(changes)
    changes["epass_seen_at"] = ts
    if substantive:
        changes["updated_at"] = ts
    if not job.get("in_feed"):
        flags = set(f for f in (job.get("flags") or "").split(",") if f)
        flags = (flags | {"parked"}) if parked else (flags - {"parked"})
        if ",".join(sorted(flags)) != (job.get("flags") or ""):
            changes["flags"] = ",".join(sorted(flags)) or None
    db.update("job", {"job_id": job["job_id"]}, changes)
    upsert_unit(db, job["job_id"], {}, model=model, serial=serial, brand=brand)
    if attached:
        return job["job_id"], "attached", new_status, parked
    return job["job_id"], ("updated" if substantive else "unchanged"), new_status, parked


def _last_first_to_first_last(name) -> str:
    """Invoice export writes 'HUCKABY NEAL & ELAINE'; DispatchTrack writes 'NEAL & ELAINE HUCKABY'. Normalise to the latter."""
    n = str(name or "").strip()
    parts = n.split(" ", 1)
    return f"{parts[1]} {parts[0]}" if len(parts) == 2 else n
